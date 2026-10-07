#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
book_converter.py -- turn a source document into a Swadhyay library book-HTML
file (and optionally register it in swadhyay-data.js / swadhyay-app.js), so
books no longer need a bespoke one-off script each time.

Supported inputs
-----------------
1. .docx  -- a Word document using paragraph STYLES to mark structure
   (Title, Author, Center, Heading 1-4, Quote, and any custom "Shlok"-named
   style). This is the preferred source: if you format the manuscript
   yourself in Word with these styles, conversion is close to 1:1.
   Real Word footnotes (Insert > Footnote) are carried over: each citation point
   becomes a superscript "[१]" marker and the note text is placed right after the
   paragraph that cites it (Footnotes, with "- source" lines as Footnotes-Right).
   A note cited from a heading goes directly under that heading, without a marker.

2. .html / .htm -- a Zoho Writer / "Scroll" knowledge-base article export
   ("Extracted Content"). Real chapter titles live in
   <div class="z-fs:24 z-fw:700 z-txt:grey0 ..."> TITLE </div> markers (NOT
   in the Aspose-style <p style="font-size:...pt"> body text) -- the tool
   segments the document on those markers.

Output format
-------------
Produces a book HTML file matching the rest of the library's established
template (see any existing file under Gita Press Books\\Gita Press Books\\
for reference): a DOCTYPE/head linking ../../master.css, body made of
<section class="epub-page" data-page="N"><div> ... </div></section> blocks
separated by <hr class="page-break">, one section per chapter. Each chapter
starts with a <p class="Heading" id="toc_marker-N"> paragraph (this is what
the in-app reader's chapter-splitter and TOC look for) followed by content
paragraphs classed from this fixed vocabulary (all styled by master.css):

    Heading              top-level chapter title -> starts a new reader page
    Sub-Heading / -2 / -3  nested chapter title (e.g. an अध्याय inside a
                         Heading-level part) -> ALSO starts its own new
                         reader page/section (independently navigable), but
                         is grouped as a child under its nearest preceding
                         Heading/Sub-Heading-N in the book-detail TOC widget
    TXT                  regular justified body prose
    Shlok                a verse / root-sutra line (red, bold, centered)
    Shlok-Right          a verse's "- source" attribution line (red, italic,
                         right-aligned)
    Footnotes            real scholarly-commentary footnote text (bordered,
                         BlueViolet) -- reserved for genuine prose footnotes,
                         not short citations (those stay Shlok/Shlok-Right)
    Footnotes-Right      a footnote's "- source" attribution line

Classification rules (used for free-text, i.e. docx "Normal" paragraphs or
any HTML paragraph that isn't an explicit chapter/sub-heading marker):
    - starts with -, En dash (U+2013) or Em dash (U+2014)   -> *-Right variant
    - matches ^\\[\\s*<devanagari-numeral>\\s*\\]            -> footnote-ish
    - everything else                                        -> TXT
  For the HTML/Zoho path specifically, a run of paragraphs starting at a
  "[N]" bracket is further split: if it contains any non-dash paragraph
  longer than --footnote-min-chars (default 200), the WHOLE run is treated
  as genuine Footnotes; otherwise it's a short verse+citation pair and
  stays Shlok/Shlok-Right. (See the भक्तियोग / राजयोग conversions this tool
  was built from for worked examples of exactly this distinction.)

Two subcommands
---------------
  convert   source .docx/.html -> book HTML + swadhyay-data.js / NON_GITA_PRESS_BOOKS entry
  cover     copy a cover image into assets/covers/ + update COVER_OVERRIDES
            (kept separate from `convert` because a cover usually arrives as
            its own image file, on its own schedule, not bundled with the
            manuscript)

Run with --help (or `convert --help` / `cover --help`) for the full flag
list. By default both subcommands WRITE to the project; pass --dry-run to
only print what would happen.

Examples
--------
  # A docx you formatted yourself with Word styles:
  python book_converter.py convert "C:\\Users\\me\\Desktop\\राजयोग.docx" ^
      --title राजयोग --author "स्वामी विवेकानन्द" --category Vedant ^
      --translator "पं. सूर्यकान्त त्रिपाठी 'निराला' तथा प्रा. श्री दिनेशचन्द्र गुह" ^
      --subtitle "(पातंजल योगसूत्र, सूत्रार्थ और व्याख्यासहित)" ^
      --publisher other

  # A messy Zoho/Aspose export, with a publisher-widget block to strip and
  # chapters that need renumbering (see --strip-between / --reorder-numeric):
  python book_converter.py convert "extracted-content.html" ^
      --title भक्तियोग --author "स्वामी विवेकानन्द" --category Vedant ^
      --strip-between प्रकाशक "Buy eBook" --reorder-numeric --publisher other

  # See what a source would produce without touching any files:
  python book_converter.py convert source.docx --title X --author Y --category Vedant --dry-run

  # Set/replace a book's cover image:
  python book_converter.py cover राजयोग.html "C:\\Users\\me\\Desktop\\rajyoga-cover.jpg"

  # (docx only) see what paragraph styles a manuscript actually uses, to
  # fill in --chapter-heading-styles / --sub-heading-styles / --quote-styles:
  python book_converter.py convert source.docx --list-styles
"""
import argparse
import html as htmlmod
import json
import re
import sys
from pathlib import Path

# ---------------------------------------------------------------- defaults -

PROJECT_ROOT = Path(__file__).resolve().parent.parent
BOOKS_DIR = PROJECT_ROOT / "Gita Press Books" / "Gita Press Books"
DATA_JS = PROJECT_ROOT / "swadhyay-data.js"
APP_JS = PROJECT_ROOT / "swadhyay-app.js"
COVERS_DIR = PROJECT_ROOT / "assets" / "covers"

# CATEGORY_META keys from swadhyay-app.js, kept here only to catch typos early.
KNOWN_CATEGORIES = [
    "all", "Vedas", "Upanishad", "Vedant", "Gita", "Purans", "Upa Puran",
    "Itihasas", "Stotra evam Naamavali", "Bajans", "Pravachan",
    "Siksha evam Katha", "Balaupayogi", "Nitya Puja evam Karmakand",
    "Swami Sharnanand Sahitya", "Teerth Sthal",
]

DEVA_DIGITS = "०१२३४५६७८९"
DASH_RE = re.compile(r"^[\-\u2013\u2014]")
BRACKET_RE = re.compile(r"^\[\s*[०-९]+\s*\]")
NUMERIC_PREFIX_RE = re.compile(r"^([०-९]+)\.")

BOOK_TEMPLATE_HEAD = """<!DOCTYPE html>
<html lang="hi">
<head>
<meta charset="UTF-8">

<title>{title} : प्रथम पृष्ठ</title>

<link rel="stylesheet" href="../../master.css">
<style>
  body {{
    font-family: Georgia, serif;
    max-width: 800px;
    margin: 2rem auto;
    padding: 0 1rem;
    line-height: 1.6;
  }}

  .page-break {{
    margin: 2rem 0;
    border: none;
    border-top: 1px dashed #ccc;
  }}

  .epub-page {{
    margin-bottom: 1rem;
  }}
</style>

</head>

<body>

"""


def esc(s: str) -> str:
    return htmlmod.escape(s, quote=True)


def deva_num(s: str):
    """Leading 'N.' in Devanagari numerals -> int, or None."""
    m = NUMERIC_PREFIX_RE.match(s)
    if not m:
        return None
    return int("".join(str(DEVA_DIGITS.index(c)) for c in m.group(1)))


class Para:
    # `html=True` means `text` is already finished HTML (an image block or a table) and must
    # be written out as-is instead of being escaped.
    __slots__ = ("text", "bold", "align", "html")

    def __init__(self, text, bold=False, align=None, html=False):
        self.text = text
        self.bold = bold
        self.align = align
        self.html = html


class Chapter:
    def __init__(self, title, css_class="Heading"):
        self.title = title
        self.css_class = css_class  # Heading / Sub-Heading / Sub-Heading-2 / Sub-Heading-3 ...
        self.paras_with_class = []  # list of (Para, class_name)

    def add(self, para: Para, cls: str):
        self.paras_with_class.append((para, cls))


# ---------------------------------------------------------- free-text rule -

def classify_free_text(text: str) -> str:
    """Fallback classification for plain prose paragraphs (docx "Normal", or
    any HTML paragraph with no more specific signal)."""
    t = text.strip()
    if DASH_RE.match(t):
        return "Shlok-Right"
    if BRACKET_RE.match(t):
        return "Footnotes"
    return "TXT"


# ===================================================== .docx conversion ===

# Word footnotes: python-docx's Paragraph.text silently drops <w:footnoteReference>, so the
# reference points are re-inserted as private-use marker characters (NOTE_OPEN <word id>
# NOTE_CLOSE) while reading, then turned into a "[१]"-style superscript at render time. The
# note text itself lives in word/footnotes.xml and is read separately.
NOTE_OPEN, NOTE_CLOSE = "", ""
NOTE_RE = re.compile(NOTE_OPEN + r"(\d+)" + NOTE_CLOSE)
_W_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def to_deva(n: int) -> str:
    return "".join(DEVA_DIGITS[int(d)] for d in str(n))


def read_docx_footnotes(path: Path):
    """{word footnote id (str): [paragraph text, ...]} from word/footnotes.xml ({} if none)."""
    import zipfile
    from xml.etree import ElementTree as ET
    try:
        z = zipfile.ZipFile(str(path))
        if "word/footnotes.xml" not in z.namelist():
            return {}
        root = ET.fromstring(z.read("word/footnotes.xml"))
    except Exception:
        return {}
    notes = {}
    for fn in root.findall(_W_NS + "footnote"):
        if fn.get(_W_NS + "type") in ("separator", "continuationSeparator", "continuationNotice"):
            continue
        paras = []
        for p in fn.iter(_W_NS + "p"):
            bits = []
            for e in p.iter():
                if e.tag == _W_NS + "t":
                    bits.append(e.text or "")
                elif e.tag == _W_NS + "tab":
                    bits.append(" ")
            t = re.sub(r"\s+", " ", "".join(bits)).strip()
            if t:
                paras.append(t)
        if paras:
            notes[fn.get(_W_NS + "id")] = paras
    return notes


def paragraph_text_with_notes(p):
    """Like python-docx's Paragraph.text, but leaves a marker where a footnote is cited."""
    from docx.text.run import Run
    from docx.oxml.ns import qn
    out = []
    for r in p._p.xpath(".//w:r"):
        for ref in r.xpath("w:footnoteReference"):
            out.append(NOTE_OPEN + ref.get(qn("w:id")) + NOTE_CLOSE)
        out.append(Run(r, p).text)
    return "".join(out)


def render_text(text: str) -> str:
    """HTML-escape paragraph text and turn footnote markers (already renumbered to their
    display number) into a superscript [१] marker."""
    return NOTE_RE.sub(lambda m: f"<sup>[{to_deva(int(m.group(1)))}]</sup>", esc(text))


def image_block_html(src: str, alt: str) -> str:
    return (f'<p class="TXT" style="text-align:center;text-indent:0;margin:1em 0">'
            f'<img src="{esc(src)}" alt="{esc(alt)}" style="max-width:100%;height:auto"></p>')


def image_row_html(srcs, alt: str) -> str:
    """Several pictures in one row: side by side when there is room (desktop), wrapping into
    a single column when there isn't (phones). Each picture asks for at least 300 px, so two
    only share a line when the reading column is wide enough for both -- no media query needed."""
    imgs = "".join(
        f'<img src="{esc(src)}" alt="{esc(alt)}" '
        f'style="flex:1 1 300px;min-width:0;max-width:100%;height:auto">' for src in srcs)
    return (f'<div style="display:flex;flex-wrap:wrap;gap:12px;justify-content:center;'
            f'align-items:flex-start;margin:1em 0">{imgs}</div>')


def table_html(rows) -> str:
    """rows: [[cell text, ...], ...] -> a bordered table whose text is set in the book's
    own body style (each cell holds a TXT paragraph, so it picks up master.css fonts).
    The first row is treated as the header."""
    def cell(tag, text, extra=""):
        inner = "<br>".join(esc(t) for t in text.split("\n") if t.strip())
        if tag == "th":
            inner = f"<b>{inner}</b>"
        return (f'<{tag} style="border:1px solid #b9a77a;padding:6px 10px;vertical-align:top;{extra}">'
                f'<p class="TXT" style="text-indent:0;text-align:left;margin:0">{inner}</p></{tag}>')
    out = ['<table style="border-collapse:collapse;margin:1em auto;width:92%">']
    for i, row in enumerate(rows):
        tag = "th" if i == 0 else "td"
        extra = "background:#f6efdc;" if i == 0 else ""
        out.append("<tr>" + "".join(cell(tag, c, extra) for c in row) + "</tr>")
    out.append("</table>")
    return "".join(out)


MAX_IMAGE_BYTES = 200_000


def optimize_image(blob: bytes, ext: str, max_bytes: int = MAX_IMAGE_BYTES):
    """Keep every picture at or under `max_bytes` (default 200 KB) with as little visible loss
    as possible. Word keeps originals (some are 3+ MB PNG scans), so anything larger is
    re-saved as JPEG: try full size (max 1600 px wide) at high quality first, lowering quality
    only a notch at a time, and shrinking the picture only if quality alone can't get there.
    Pictures already within the limit are kept byte-for-byte."""
    if len(blob) <= max_bytes:
        return blob, ext
    try:
        import io
        from PIL import Image
        im = Image.open(io.BytesIO(blob))
        if im.mode in ("RGBA", "LA", "P"):
            im = im.convert("RGBA")
            bg = Image.new("RGB", im.size, "white")
            bg.paste(im, mask=im.split()[-1])
            im = bg
        else:
            im = im.convert("RGB")
        if im.width > 1600:
            im = im.resize((1600, round(im.height * 1600 / im.width)), Image.LANCZOS)
        best = None
        for scale in (1.0, 0.92, 0.85, 0.78, 0.7, 0.62, 0.55, 0.48):
            work = im if scale == 1.0 else im.resize(
                (max(1, round(im.width * scale)), max(1, round(im.height * scale))), Image.LANCZOS)
            for q in (92, 89, 86, 83, 80, 77, 74):
                out = io.BytesIO()
                work.save(out, "JPEG", quality=q, optimize=True, progressive=True)
                data = out.getvalue()
                if best is None or len(data) < len(best):
                    best = data
                if len(data) <= max_bytes:
                    return data, ".jpg"
        return best, ".jpg"
    except Exception:
        return blob, ext


def read_docx_items(doc, image_slug: str, image_store: list):
    """Walk the document body IN ORDER and return a list of items
    {"style", "text", "images": [filename, ...], "table": [[cell, ...], ...] or None}.
    Paragraph text keeps footnote markers; embedded pictures are recorded in `image_store`
    as (filename, bytes) -- written to the book's image/ folder by cmd_convert -- and
    tables are kept (python-docx's `doc.paragraphs` skips both)."""
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    from docx.oxml.ns import qn
    items = []
    for el in doc.element.body.iterchildren():
        if el.tag == qn("w:p"):
            p = Paragraph(el, doc)
            names = []
            for blip in el.xpath(".//a:blip"):
                rid = blip.get(qn("r:embed"))
                if not rid or rid not in doc.part.related_parts:
                    continue
                part = doc.part.related_parts[rid]
                ext = part.partname.ext.lower()
                ext = ext if ext.startswith(".") else "." + ext
                blob, ext = optimize_image(part.blob, ext)
                name = f"{image_slug}-{len(image_store) + 1}{ext}"
                image_store.append((name, blob))
                names.append(name)
            items.append({"style": p.style.name, "text": paragraph_text_with_notes(p).strip(),
                          "images": names, "table": None})
        elif el.tag == qn("w:tbl"):
            t = Table(el, doc)
            rows = [[c.text.strip() for c in r.cells] for r in t.rows]
            if any(any(c for c in r) for r in rows):
                items.append({"style": "__table__", "text": "", "images": [], "table": rows})
    keep = lambda it: (it["text"].replace(NOTE_OPEN, "").replace(NOTE_CLOSE, "").strip()
                       or it["images"] or it["table"])
    return [it for it in items if keep(it)]


def convert_docx(path: Path, args):
    try:
        import docx
    except ImportError:
        sys.exit("python-docx is required for .docx input: pip install python-docx")

    doc = docx.Document(str(path))
    footnotes = read_docx_footnotes(path)
    image_slug = devanagari_to_slug(args.out_name or args.title or "book")
    image_store = []
    items = read_docx_items(doc, image_slug, image_store)
    # `raw` keeps the (style, text) shape the rest of this function was written around;
    # `extras[i]` carries that paragraph's pictures / table.
    raw = [(it["style"], it["text"]) for it in items]
    extras = [(it["images"], it["table"]) for it in items]

    chapter_styles = set(s.strip() for s in args.chapter_heading_styles.split(","))
    sub_styles = [s.strip() for s in args.sub_heading_styles.split(",") if s.strip()]
    # Every heading-ish style -- chapter-level AND sub-level -- starts its own
    # reader-navigable <section> (page-break boundary); the CSS class assigned
    # here only controls how the book-detail TOC widget nests/groups them
    # (extractTOC/groupTOC in swadhyay-app.js walk the resulting classes, not
    # the section boundaries). First sub-style -> Sub-Heading, next -> Sub-
    # Heading-2, next -> Sub-Heading-3, matching master.css's nesting classes.
    SUB_HEADING_CLASSES = ["Sub-Heading", "Sub-Heading-2", "Sub-Heading-3", "Sub-Heading-4"]
    heading_style_class = {s: "Heading" for s in chapter_styles}
    for i, s in enumerate(sub_styles):
        heading_style_class[s] = SUB_HEADING_CLASSES[min(i, len(SUB_HEADING_CLASSES) - 1)]

    quote_like_styles = set(s.strip() for s in args.quote_styles.split(","))

    # ---- title-page block: everything before the first chapter heading ----
    try:
        title_end = next(i for i, (s, t) in enumerate(raw) if s in heading_style_class)
    except StopIteration:
        sys.exit(f"No paragraph uses any of --chapter-heading-styles ({sorted(chapter_styles)}); "
                  f"nothing to split into chapters. Check the style names with --list-styles.")

    title_block = raw[:title_end]
    body_block = raw[title_end:]
    title_extras = extras[:title_end]
    body_extras = extras[title_end:]
    # pictures that sit on the title page (before the first heading) go on the cover
    cover_images = [n for imgs, _tbl in title_extras for n in imgs]

    # Safety check: the manuscript's own "Title" paragraph must be the book being imported.
    # Without this, converting the wrong .docx under the right --title silently produces a
    # book with one book's name on another book's text (this happened once: श्रीकृष्ण चैतन्य
    # was built from the राजयोग manuscript).
    def _norm_title(s):
        return re.sub(r"[\s\-–—_.,:;!?()\[\]\"'‘’“”।॥]+", "", s)
    doc_title = next((t for s, t in title_block if s == "Title"), "")
    if doc_title and _norm_title(doc_title) != _norm_title(args.title) and not args.allow_title_mismatch:
        sys.exit(f"STOP: the document's own title is '{doc_title}' but --title is '{args.title}'. "
                 f"This looks like the wrong source file. Check the file, or pass "
                 f"--allow-title-mismatch if the difference is intentional.")

    drop_titles = set(s.strip() for s in args.drop_chapters.split(",") if s.strip())
    right_styles = set(s.strip() for s in args.right_styles.split(",") if s.strip())
    center_styles = set(s.strip() for s in args.center_styles.split(",") if s.strip())

    book_title = args.title
    subtitle = args.subtitle
    author_name = args.author
    translator_line = ("अनुवादक: " + args.translator) if args.translator else ""
    # fall back to whatever the docx's own title-page paragraphs say, for any
    # field the caller didn't pass explicitly.
    if not subtitle:
        for s, t in title_block:
            if s == "Center" and t.startswith("("):
                subtitle = t
                break

    # Footnotes are numbered 1, 2, 3... in reading order (the number Word displays), not by
    # Word's internal ids, and each note's text is placed straight after the paragraph that
    # cites it so it stays on the same reader page. A note cited from a chapter heading goes
    # right under that heading instead (a marker inside the heading would end up in the TOC).
    note_no = [0]

    def renumber(text):
        def sub(m):
            note_no[0] += 1
            return NOTE_OPEN + str(note_no[0]) + NOTE_CLOSE
        return NOTE_RE.sub(sub, text)

    def note_paragraphs(word_id, number):
        out = []
        for i, t in enumerate(footnotes.get(word_id, [])):
            cls = "Footnotes-Right" if DASH_RE.match(t) else "Footnotes"
            if i == 0:
                t = f"[{to_deva(number)}] {t}"
            out.append((Para(t), cls))
        return out

    chapters = []
    current = None
    # Chapters named in --drop-chapters (e.g. a book's own contents list, which the reader's
    # TOC drawer makes redundant) are left out -- but any picture inside one is carried over to
    # the start of the next kept chapter instead of disappearing with it.
    drop_current = False
    carry = []
    carry_names = []     # file names of those carried pictures (for --carried-images-chapter)
    def image_paras(names):
        return [(Para(image_block_html("image/" + n, f"{args.title} — चित्र"), html=True), "Image") for n in names]

    for (style, text), (imgs, table) in zip(body_block, body_extras):
        word_ids = NOTE_RE.findall(text)
        first_no = note_no[0] + 1
        if word_ids:
            text = renumber(text)
        if table is not None:
            if current is not None:
                current.add(Para(table_html(table), html=True), "Table")
            continue
        if style in heading_style_class:
            if current is not None and drop_current:
                carry += [pc for pc in current.paras_with_class if pc[1] == "Image"]
            heading_text = NOTE_RE.sub("", text).strip()
            current = Chapter(heading_text, heading_style_class[style])
            drop_current = heading_text in drop_titles
            if drop_current:
                print(f"[drop] leaving out chapter: {heading_text}")
                carry_names += imgs
            else:
                if carry_names and args.carried_images_chapter:
                    # the pictures left over from dropped chapters get a chapter of their own,
                    # placed just before this one, laid out in a responsive row
                    gallery = Chapter(args.carried_images_chapter, "Heading")
                    gallery.add(Para(image_row_html(["image/" + n for n in carry_names],
                                                    f"{args.title} — चित्र"), html=True), "Image")
                    chapters.append(gallery)
                    print(f"[images] {len(carry_names)} picture(s) moved to new chapter: {args.carried_images_chapter}")
                    carry, carry_names = [], []
                chapters.append(current)
                for k, wid in enumerate(word_ids):
                    for para, cls in note_paragraphs(wid, first_no + k):
                        current.add(para, cls)
                for para, cls in carry:
                    current.add(para, cls)
                carry, carry_names = [], []
            for para, cls in image_paras(imgs):
                current.add(para, cls)
            continue
        if current is None:
            continue  # shouldn't happen: body_block starts at title_end
        if drop_current:
            carry_names += imgs
        for para, cls in image_paras(imgs):      # a picture goes just above its paragraph
            current.add(para, cls)
        if not NOTE_RE.sub("", text).strip():
            continue                              # picture-only paragraph
        if style in quote_like_styles or style == args.shlok_style:
            current.add(Para(text), "Shlok")
        elif style in right_styles:
            current.add(Para(text), "TXT-Right")      # signature / attribution lines
        elif style in center_styles:
            current.add(Para(text), "Shlok")          # centred closing lines (e.g. शान्ति-पाठ)
        else:
            current.add(Para(text), classify_free_text(NOTE_RE.sub("", text)))
        for k, wid in enumerate(word_ids):
            for para, cls in note_paragraphs(wid, first_no + k):
                current.add(para, cls)

    return {
        "book_title": book_title, "subtitle": subtitle, "author_name": author_name,
        "translator_line": translator_line, "chapters": chapters,
        "cover_images": cover_images, "images": image_store,
    }


def list_docx_styles(path: Path):
    import docx
    from collections import Counter
    doc = docx.Document(str(path))
    counts = Counter(p.style.name for p in doc.paragraphs if p.text.strip())
    print(f"Paragraph styles used in {path.name}:")
    for s, c in counts.most_common():
        print(f"  {s}: {c}")


# ================================================ Zoho/Aspose HTML conversion

def convert_zoho_html(path: Path, args):
    raw = path.read_text(encoding="utf-8")

    marker_re = re.compile(
        r'<div class="' + re.escape(args.marker_class) + r'[^"]*">\s*(.*?)\s*</div>',
        re.DOTALL,
    )
    markers = [
        (m.start(), m.end(), re.sub(r"<[^>]+>", "", m.group(1)).strip())
        for m in marker_re.finditer(raw)
    ]
    if not markers:
        sys.exit(f'No chapter markers found (looked for class="{args.marker_class}..."). '
                  f"Pass --marker-class to match this export's actual marker div class.")

    chunks = []
    for i, (s, e, title) in enumerate(markers):
        chunk_end = markers[i + 1][0] if i + 1 < len(markers) else len(raw)
        chunks.append((title, raw[e:chunk_end]))

    p_re = re.compile(r"<p\b([^>]*)>(.*?)</p>", re.DOTALL)
    span_re = re.compile(r"<span\b([^>]*)>(.*?)</span>", re.DOTALL)
    tag_re = re.compile(r"<[^>]+>")
    img_re = re.compile(r"<img\b")
    align_re = re.compile(r"text-align:\s*(\w+)")
    fw_re = re.compile(r"font-weight:\s*(\d+)")

    def clean(s):
        s = tag_re.sub("", s)
        s = htmlmod.unescape(s)
        s = s.replace("\xa0", " ")
        return re.sub(r"\s+", " ", s).strip()

    def extract_paragraphs(frag):
        frag = re.sub(r"<style\b[^>]*>.*?</style>", "", frag, flags=re.DOTALL)
        frag = re.sub(r"<svg\b[^>]*>.*?</svg>", "", frag, flags=re.DOTALL)
        out = []
        for m in p_re.finditer(frag):
            pattrs, inner = m.group(1), m.group(2)
            if img_re.search(inner):
                continue
            spans = span_re.findall(inner)
            bold = False
            if spans:
                parts = []
                for sattrs, scontent in spans:
                    fwm = fw_re.search(sattrs)
                    if fwm and int(fwm.group(1)) >= 600:
                        bold = True
                    parts.append(clean(scontent))
                text = " ".join(t for t in parts if t)
            else:
                text = clean(inner)
            if not text:
                continue
            am = align_re.search(pattrs)
            out.append(Para(text, bold=bold, align=(am.group(1) if am else None)))
        return out

    titled_chunks = [(title, extract_paragraphs(frag)) for title, frag in chunks]

    # ---- optional: strip a boilerplate run (e.g. a publisher/"Buy" widget) ----
    for start_text, end_text in args.strip_between:
        for title, paras in titled_chunks:
            out, skipping = [], False
            for p in paras:
                if not skipping and p.text == start_text:
                    skipping = True
                    continue
                if skipping:
                    if p.text == end_text:
                        skipping = False
                    continue
                out.append(p)
            paras[:] = out

    titled_chunks = [(t, ps) for t, ps in titled_chunks if ps]  # drop empty (e.g. an image-only page)

    if args.reorder_numeric:
        numbered = [(t, p) for t, p in titled_chunks if deva_num(t) is not None]
        unnumbered = [(t, p) for t, p in titled_chunks if deva_num(t) is None]
        numbered.sort(key=lambda tp: deva_num(tp[0]))
        titled_chunks = unnumbered + numbered

    chapters = []
    for title, paras in titled_chunks:
        ch = Chapter(title)
        bracket_idxs = [i for i, p in enumerate(paras) if BRACKET_RE.match(p.text)]
        zone_start = bracket_idxs[0] if bracket_idxs else len(paras)
        for i in range(zone_start):
            ch.add(paras[i], classify_free_text(paras[i].text))
        if bracket_idxs:
            for ri, b in enumerate(bracket_idxs):
                run_end = bracket_idxs[ri + 1] if ri + 1 < len(bracket_idxs) else len(paras)
                run = paras[b:run_end]
                non_dash = [p.text for p in run if not DASH_RE.match(p.text)]
                is_footnote = any(len(t) > args.footnote_min_chars for t in non_dash)
                for p in run:
                    if DASH_RE.match(p.text):
                        ch.add(p, "Footnotes-Right" if is_footnote else "Shlok-Right")
                    else:
                        ch.add(p, "Footnotes" if is_footnote else "Shlok")
        chapters.append(ch)

    return {
        "book_title": args.title, "subtitle": args.subtitle, "author_name": args.author,
        "translator_line": ("अनुवादक: " + args.translator) if args.translator else "",
        "chapters": chapters,
    }


# ============================================================ HTML output =

def build_html(result, args) -> str:
    title = result["book_title"]
    sections = ['<section class="epub-page" data-page="1">\n<div class="Basic-Graphics-Frame">\n'
                f'\t\t\t<p class="Title-Page---Book-Name">{esc(title)}</p>']
    if result["author_name"]:
        sections.append(f'\t\t\t<p class="Title-Page---Author-Name">{esc(result["author_name"])}</p>')
    if result["translator_line"]:
        sections.append(f'\t\t\t<p class="Title-Page---Translater-Author-Name">{esc(result["translator_line"])}</p>')
    for name in result.get("cover_images", []):
        sections.append("\t\t\t" + image_block_html("image/" + name, f"{title} — मुखपृष्ठ चित्र"))
    if result["subtitle"]:
        sections.append(f'\t\t\t<p class="Title-Page---Publisher-Note">{esc(result["subtitle"])}</p>')
    sections.append("\t\t</div>\n</section>")
    title_section = "\n".join(sections)

    body_sections = []
    toc_counter = 0
    for i, ch in enumerate(result["chapters"]):
        toc_counter += 1
        lines = [f'\t\t\t<p id="toc_marker-{toc_counter}" class="{ch.css_class}">{esc(ch.title)}</p>']
        for para, cls in ch.paras_with_class:
            if para.html:
                lines.append("\t\t\t" + para.text)
            else:
                lines.append(f'\t\t\t<p class="{cls}">{render_text(para.text)}</p>')
        body_sections.append(
            f'<section class="epub-page" data-page="{i + 2}">\n<div>\n' + "\n".join(lines) + "\n\t\t</div>\n</section>"
        )

    all_sections = [title_section] + body_sections
    body_html = "\n\n<hr class=\"page-break\">\n\n".join(all_sections)
    return BOOK_TEMPLATE_HEAD.format(title=esc(title)) + body_html + "\n\n</body>\n</html>\n"


def print_stats(result):
    print(f"Title: {result['book_title']}")
    print(f"Author: {result['author_name']}")
    if result["translator_line"]:
        print(f"Translator line: {result['translator_line']}")
    if result["subtitle"]:
        print(f"Subtitle: {result['subtitle']}")
    print(f"Chapters: {len(result['chapters'])}")
    if result.get("images"):
        print(f"Pictures: {len(result['images'])}  (cover: {len(result.get('cover_images', []))})")
    stats = {}
    for ch in result["chapters"]:
        for _, cls in ch.paras_with_class:
            stats[cls] = stats.get(cls, 0) + 1
    print("Paragraph classification counts:", stats)
    print("\nChapter list:")
    for i, ch in enumerate(result["chapters"]):
        print(f"  {i + 1}. {ch.title}  ({len(ch.paras_with_class)} paragraphs)")


# ================================================= project file registration

def register_in_data_js(data_js: Path, filename: str, category: str, author: str, size_kb: float, dry_run: bool):
    text = data_js.read_text(encoding="utf-8")
    if f'"file":  "{filename}"' in text or f'"file": "{filename}"' in text:
        print(f"[data.js] {filename} already registered -- skipping (edit swadhyay-data.js by hand to change it).")
        return
    entry = (
        '    {\n'
        f'        "file":  "{filename}",\n'
        f'        "category":  "{category}",\n'
        f'        "author":  "{author}",\n'
        '        "tikakar":  "",\n'
        '        "translator":  "",\n'
        '        "publisher":  "",\n'
        f'        "sizeKB":  {size_kb}\n'
        '    }'
    )
    # The array's closing "]" is the last one in the file (it's followed only by
    # an optional ";" and whitespace) -- don't require it to be the very last
    # character, since some editors/tools append "];" or a trailing newline.
    idx = text.rfind("]")
    if idx == -1:
        sys.exit(f"Could not find the closing ']' of SWADHYAY_BOOKS_RAW in {data_js}; add the entry by hand:\n{entry}\n]")
    prefix = text[:idx].rstrip()
    if prefix.endswith("}"):
        prefix += ","
    new_text = prefix + "\n" + entry + "\n]" + text[idx + 1:]
    if dry_run:
        print(f"[dry-run] would append to {data_js.name}:\n{entry}")
    else:
        data_js.write_text(new_text, encoding="utf-8")
        print(f"[data.js] registered {filename} ({category}, {author}, {size_kb} KB)")


def register_non_gita_press(app_js: Path, filename: str, dry_run: bool):
    text = app_js.read_text(encoding="utf-8")
    if f'"{filename}": true' in text:
        print(f"[app.js] {filename} already in NON_GITA_PRESS_BOOKS -- skipping.")
        return
    marker = "var NON_GITA_PRESS_BOOKS = {\n"
    idx = text.find(marker)
    if idx == -1:
        sys.exit(f"Could not find NON_GITA_PRESS_BOOKS in {app_js}; add by hand: \"{filename}\": true,")
    insert_at = idx + len(marker)
    new_text = text[:insert_at] + f'    "{filename}": true,\n' + text[insert_at:]
    if dry_run:
        print(f'[dry-run] would add "{filename}": true to NON_GITA_PRESS_BOOKS in {app_js.name}')
    else:
        app_js.write_text(new_text, encoding="utf-8")
        print(f"[app.js] added {filename} to NON_GITA_PRESS_BOOKS")


# ============================================================= cover images

# Rough Devanagari -> ASCII transliteration, good enough for a readable slug
# (not meant to be linguistically precise -- pass --slug to override it).
_DEVA_INDEPENDENT_VOWELS = {
    "अ": "a", "आ": "aa", "इ": "i", "ई": "ee", "उ": "u", "ऊ": "oo",
    "ऋ": "ri", "ए": "e", "ऐ": "ai", "ओ": "o", "औ": "au", "ं": "n", "ः": "h", "ँ": "n",
}
_DEVA_VOWEL_SIGNS = {
    "ा": "aa", "ि": "i", "ी": "ee", "ु": "u", "ू": "oo", "ृ": "ri",
    "े": "e", "ै": "ai", "ो": "o", "ौ": "au", "्": "",
}
_DEVA_CONSONANTS = {
    "क": "k", "ख": "kh", "ग": "g", "घ": "gh", "ङ": "ng",
    "च": "ch", "छ": "chh", "ज": "j", "झ": "jh", "ञ": "ny",
    "ट": "t", "ठ": "th", "ड": "d", "ढ": "dh", "ण": "n",
    "त": "t", "थ": "th", "द": "d", "ध": "dh", "न": "n",
    "प": "p", "फ": "ph", "ब": "b", "भ": "bh", "म": "m",
    "य": "y", "र": "r", "ल": "l", "व": "v",
    "श": "sh", "ष": "sh", "स": "s", "ह": "h",
    "क्ष": "ksh", "त्र": "tr", "ज्ञ": "gy", "ड़": "r", "ढ़": "rh",
}
_DEVA_DIGIT_SLUG = {d: str(i) for i, d in enumerate(DEVA_DIGITS)}


def devanagari_to_slug(text: str) -> str:
    """Best-effort romanized, hyphenated slug from a Devanagari string.
    Good enough to match the existing assets/covers/*.* naming convention;
    pass --slug explicitly if the result isn't what you want."""
    out = []
    i = 0
    n = len(text)
    while i < n:
        ch = text[i]
        two = text[i:i + 2]
        if two in _DEVA_CONSONANTS:
            out.append(_DEVA_CONSONANTS[two])
            i += 2
            continue
        if ch in _DEVA_CONSONANTS:
            cons = _DEVA_CONSONANTS[ch]
            nxt = text[i + 1] if i + 1 < n else ""
            if nxt in _DEVA_VOWEL_SIGNS:
                out.append(cons + _DEVA_VOWEL_SIGNS[nxt])
                i += 2
            elif nxt and (nxt in _DEVA_CONSONANTS or (len(text[i + 1:i + 3]) == 2 and text[i + 1:i + 3] in _DEVA_CONSONANTS)):
                out.append(cons + "a")  # inherent vowel before another consonant
                i += 1
            else:
                out.append(cons)  # virama / end-of-word: drop inherent vowel
                i += 1
            continue
        if ch in _DEVA_INDEPENDENT_VOWELS:
            out.append(_DEVA_INDEPENDENT_VOWELS[ch])
            i += 1
            continue
        if ch in _DEVA_DIGIT_SLUG:
            out.append(_DEVA_DIGIT_SLUG[ch])
            i += 1
            continue
        if ch.isalnum():
            out.append(ch.lower())
            i += 1
            continue
        out.append(" ")
        i += 1
    slug = "".join(out)
    slug = re.sub(r"[^a-z0-9]+", "-", slug.lower()).strip("-")
    slug = re.sub(r"-{2,}", "-", slug)
    return slug or "cover"


def update_cover_override(app_js: Path, filename: str, cover_rel_path: str, dry_run: bool):
    text = app_js.read_text(encoding="utf-8")
    entry_re = re.compile(r'(\n\s*"' + re.escape(filename) + r'"\s*:\s*)"[^"]*"(,?)')
    new_line = f'    "{filename}": "{cover_rel_path}",'
    if entry_re.search(text):
        new_text = entry_re.sub(lambda m: m.group(1) + f'"{cover_rel_path}"' + m.group(2), text)
        action = "updated"
    else:
        marker = "var COVER_OVERRIDES = {\n"
        idx = text.find(marker)
        if idx == -1:
            sys.exit(f"Could not find COVER_OVERRIDES in {app_js}; add by hand:\n{new_line}")
        insert_at = idx + len(marker)
        new_text = text[:insert_at] + new_line + "\n" + text[insert_at:]
        action = "added"
    if dry_run:
        print(f"[dry-run] would have {action} COVER_OVERRIDES[\"{filename}\"] = \"{cover_rel_path}\" in {app_js.name}")
    else:
        app_js.write_text(new_text, encoding="utf-8")
        print(f"[app.js] {action} COVER_OVERRIDES[\"{filename}\"] = \"{cover_rel_path}\"")


def cmd_set_cover(args):
    image_src: Path = args.image
    if not image_src.exists():
        sys.exit(f"Cover image not found: {image_src}")
    ext = image_src.suffix.lower()
    if ext not in (".png", ".jpg", ".jpeg", ".webp"):
        sys.exit(f"Unsupported cover image type {ext}; use .png, .jpg/.jpeg or .webp")

    book_file = args.book_file
    if not book_file.endswith(".html"):
        book_file += ".html"

    slug = args.slug or devanagari_to_slug(book_file[:-5])
    dest_name = f"{slug}{ext}"
    dest_path = args.covers_dir / dest_name
    rel_path = f"assets/covers/{dest_name}"

    print(f"Book file: {book_file}")
    print(f"Cover source: {image_src}  ({image_src.stat().st_size / 1024:.1f} KB)")
    print(f"Cover destination: {dest_path}")
    print(f"COVER_OVERRIDES value: {rel_path}")

    if args.dry_run:
        print("[dry-run] not copying the image or editing swadhyay-app.js.")
        return

    args.covers_dir.mkdir(parents=True, exist_ok=True)
    if dest_path.exists() and not args.force:
        sys.exit(f"{dest_path} already exists -- pass --force to overwrite, or --slug to pick a different name.")
    dest_path.write_bytes(image_src.read_bytes())
    print(f"Copied cover to {dest_path}")

    update_cover_override(args.app_js, book_file, rel_path, dry_run=False)


# ===================================================================== CLI =

def cmd_convert(args):
    if not args.source.exists():
        sys.exit(f"Source not found: {args.source}")

    suffix = args.source.suffix.lower()
    if args.list_styles:
        if suffix != ".docx":
            sys.exit("--list-styles only applies to .docx sources")
        list_docx_styles(args.source)
        return

    missing = [n for n, v in (("--title", args.title), ("--author", args.author), ("--category", args.category)) if not v]
    if missing:
        sys.exit(f"Missing required argument(s): {', '.join(missing)}")

    if suffix == ".docx":
        result = convert_docx(args.source, args)
    elif suffix in (".html", ".htm"):
        result = convert_zoho_html(args.source, args)
    else:
        sys.exit(f"Unsupported source type: {suffix} (expected .docx, .html or .htm)")

    print_stats(result)

    out_name = (args.out_name or args.title) + ".html"
    out_path = args.out_dir / out_name
    html_out = build_html(result, args)
    size_kb = round(len(html_out.encode("utf-8")) / 1024, 1)

    print(f"\nOutput: {out_path}  ({size_kb} KB)")
    if args.dry_run:
        print("[dry-run] not writing any files.")
    else:
        args.out_dir.mkdir(parents=True, exist_ok=True)
        out_path.write_text(html_out, encoding="utf-8")
        print(f"Wrote {out_path}")
        images = result.get("images") or []
        if images:
            img_dir = args.out_dir / "image"
            img_dir.mkdir(parents=True, exist_ok=True)
            for name, blob in images:
                (img_dir / name).write_bytes(blob)
            print(f"Wrote {len(images)} image(s) to {img_dir}")

    if not args.no_register:
        register_in_data_js(args.data_js, out_name, args.category, args.author, size_kb, args.dry_run)
        if args.publisher == "other":
            register_non_gita_press(args.app_js, out_name, args.dry_run)


def build_arg_parser():
    ap = argparse.ArgumentParser(
        prog="book_converter.py",
        description="Import books into the Swadhyay library from .docx/.html, and manage their cover images.",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog=__doc__,
    )
    sub = ap.add_subparsers(dest="command", required=True)

    cv = sub.add_parser("convert", help="convert a .docx or Zoho/Aspose .html source into a book-HTML file",
                         formatter_class=argparse.RawDescriptionHelpFormatter, epilog=__doc__)
    cv.add_argument("source", type=Path, help="source .docx or .html file")
    cv.add_argument("--title", help="book title (required, unless --list-styles)")
    cv.add_argument("--author", help="author name shown in the library (required, unless --list-styles)")
    cv.add_argument("--category", choices=KNOWN_CATEGORIES, help="CATEGORY_META key from swadhyay-app.js (required, unless --list-styles)")
    cv.add_argument("--translator", default="", help="translator credit line, without the 'अनुवादक:' prefix (added automatically)")
    cv.add_argument("--subtitle", default="", help="small title-page note, e.g. '(पातंजल योगसूत्र, ...)'")
    cv.add_argument("--publisher", choices=["gitapress", "other"], default="gitapress",
                     help="'other' also registers the file in NON_GITA_PRESS_BOOKS and tunes the book-detail blurb")
    cv.add_argument("--out-dir", type=Path, default=BOOKS_DIR, help="where to write <title>.html (default: the project's book folder)")
    cv.add_argument("--out-name", default=None, help="output filename without extension (default: --title)")
    cv.add_argument("--data-js", type=Path, default=DATA_JS, help="path to swadhyay-data.js")
    cv.add_argument("--app-js", type=Path, default=APP_JS, help="path to swadhyay-app.js")
    cv.add_argument("--no-register", action="store_true", help="only write the book HTML; don't touch swadhyay-data.js / swadhyay-app.js")
    cv.add_argument("--dry-run", action="store_true", help="print stats and what would be written/registered, but don't write anything")
    cv.add_argument("--list-styles", action="store_true", help="(docx only) print paragraph style usage counts and exit")

    docx_group = cv.add_argument_group(".docx options")
    docx_group.add_argument("--chapter-heading-styles", default="Heading 1,Heading 2",
                             help="comma-separated Word style names that each start a new reader chapter")
    docx_group.add_argument("--sub-heading-styles", default="Heading 3,Heading 4",
                             help="comma-separated Word style names for nested chapters -- each still "
                                  "starts its own reader page, but nests under its parent in the TOC "
                                  "(1st listed style -> Sub-Heading, 2nd -> Sub-Heading-2, 3rd -> Sub-Heading-3, ...)")
    docx_group.add_argument("--quote-styles", default="Quote",
                             help="comma-separated Word style names to render as a display Shlok block")
    docx_group.add_argument("--shlok-style", default="Shlok",
                             help="Word style name (if any) the source already uses for verse/sutra lines")
    docx_group.add_argument("--right-styles", default="right-text",
                             help="comma-separated Word style names for right-aligned lines (signatures, "
                                  "attributions) -> TXT-Right")
    docx_group.add_argument("--center-styles", default="center-text",
                             help="comma-separated Word style names for centred body lines -> Shlok")
    docx_group.add_argument("--drop-chapters", default="",
                             help="comma-separated exact heading texts to leave out (e.g. the book's own "
                                  "contents list, 'विषय क्रम'); pictures inside them move to the next chapter")
    docx_group.add_argument("--carried-images-chapter", default="",
                             help="title for a NEW chapter that collects the pictures left over from "
                                  "--drop-chapters (laid out side by side on wide screens, stacked on "
                                  "phones); without it they join the start of the next chapter")
    docx_group.add_argument("--allow-title-mismatch", action="store_true",
                             help="don't stop when the document's own Title paragraph differs from --title")

    html_group = cv.add_argument_group(".html (Zoho/Aspose export) options")
    html_group.add_argument("--marker-class", default="z-fs:24 z-fw:700 z-txt:grey0",
                             help="the chapter-title marker div's class prefix (see module docstring)")
    html_group.add_argument("--strip-between", nargs=2, metavar=("START_TEXT", "END_TEXT"), action="append", default=[],
                             help="drop every paragraph from one whose text exactly equals START_TEXT through one "
                                  "equal to END_TEXT (inclusive); repeatable. Use for publisher/\"Buy\" widgets etc.")
    html_group.add_argument("--reorder-numeric", action="store_true",
                             help="sort chapters whose titles start with 'N.' (Devanagari numerals) into numeric "
                                  "order after the unnumbered front matter, fixing any out-of-order export")
    html_group.add_argument("--footnote-min-chars", type=int, default=200,
                             help="a '[N]...' citation run becomes Footnotes (not Shlok) if any line in it exceeds this length")
    cv.set_defaults(func=cmd_convert)

    cov = sub.add_parser("cover", help="set/replace a book's cover image (separate from convert -- "
                                        "covers usually arrive on their own, after the text)")
    cov.add_argument("book_file", help="the book's registered filename, e.g. राजयोग or राजयोग.html")
    cov.add_argument("image", type=Path, help="cover image file (.png, .jpg/.jpeg or .webp)")
    cov.add_argument("--slug", default=None,
                      help="filename (without extension) to save under in assets/covers/ "
                           "(default: auto-romanized from the book filename)")
    cov.add_argument("--covers-dir", type=Path, default=COVERS_DIR, help="covers directory (default: assets/covers)")
    cov.add_argument("--app-js", type=Path, default=APP_JS, help="path to swadhyay-app.js")
    cov.add_argument("--force", action="store_true", help="overwrite the destination image file if it already exists")
    cov.add_argument("--dry-run", action="store_true", help="print what would happen, but don't write anything")
    cov.set_defaults(func=cmd_set_cover)

    return ap


def main():
    args = build_arg_parser().parse_args()
    args.func(args)


if __name__ == "__main__":
    main()
